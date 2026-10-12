"""Loading and chunking of uploaded files (code, PDF, Word, PowerPoint, photos) with LangChain.

PDFs are read page by page with PyMuPDF, slides one by one, so every chunk keeps the page/slide
it came from. Scanned PDF pages and photos have no text of their own; they are read by the vision
model (read_page). Chunks are cut on natural boundaries (paragraphs, sentences, or functions/classes
for source code) instead of at a fixed character count.
"""
import re
from typing import Callable

import pymupdf
from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_text_splitters import Language, RecursiveCharacterTextSplitter
from pptx import Presentation

CHUNK_SIZE = 1200
CHUNK_OVERLAP = 200

TEXT_EXTENSIONS = {
    "txt", "md", "rst", "csv", "json", "xml", "yaml", "yml", "toml", "ini", "cfg", "env", "log", "sql",
    "py", "js", "jsx", "ts", "tsx", "java", "kt", "c", "h", "cpp", "hpp", "cs", "go", "rs", "rb", "php",
    "swift", "scala", "dart", "r", "sh", "bat", "ps1", "html", "htm", "css", "scss", "vue", "svelte", "ipynb",
}
IMAGE_TYPES = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "webp": "image/webp"}
ALLOWED_EXTENSIONS = TEXT_EXTENSIONS | {"pdf", "docx", "pptx"} | set(IMAGE_TYPES)
# A PDF page with a picture and fewer characters than this is a scan
SCANNED_PAGE_CHARS = 20
SCAN_DPI = 150

# Extensions that get a language-aware splitter (keeps functions and classes together)
CODE_LANGUAGES = {
    "py": Language.PYTHON, "js": Language.JS, "jsx": Language.JS, "ts": Language.TS, "tsx": Language.TS,
    "java": Language.JAVA, "kt": Language.KOTLIN, "c": Language.C, "h": Language.C, "cpp": Language.CPP,
    "hpp": Language.CPP, "cs": Language.CSHARP, "go": Language.GO, "rs": Language.RUST, "rb": Language.RUBY,
    "php": Language.PHP, "swift": Language.SWIFT, "scala": Language.SCALA, "html": Language.HTML,
    "htm": Language.HTML, "md": Language.MARKDOWN, "rst": Language.RST,
}


def file_extension(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _clean(text: str) -> str:
    """Tidy text extracted from PDFs: re-join hyphenated line breaks and collapse stray whitespace."""
    text = text.replace("\x00", "")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _table_rows(table) -> list[str]:
    return [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]


def _load_pdf(file_path: str, read_page: Callable[[bytes, str], str], max_scanned: int) -> list[Document]:
    with pymupdf.open(file_path) as pdf:
        texts = [_clean(page.get_text()) for page in pdf]
        scanned = [
            i for i, page in enumerate(pdf) if len(texts[i]) < SCANNED_PAGE_CHARS and page.get_images()
        ]
        if len(scanned) > max_scanned:
            raise ValueError(
                f"{len(scanned)} scanned pages, and at most {max_scanned} can be read per file. "
                "Split the file into smaller parts."
            )
        for i in scanned:
            texts[i] = read_page(pdf[i].get_pixmap(dpi=SCAN_DPI).tobytes("png"), "image/png")
    return [Document(page_content=text, metadata={"page": number}) for number, text in enumerate(texts, start=1)]


def load_documents(
    file_path: str, ext: str, read_page: Callable[[bytes, str], str], max_scanned: int
) -> list[Document]:
    """One Document per PDF page / slide (with its number in metadata), or one for the whole file.

    read_page(image bytes, mime type) returns the text of a scanned page or a photo.
    Raises ValueError with a user-facing reason when the file has more scanned pages than max_scanned.
    """
    if ext == "pdf":
        return _load_pdf(file_path, read_page, max_scanned)

    if ext in IMAGE_TYPES:
        with open(file_path, "rb") as f:
            return [Document(page_content=read_page(f.read(), IMAGE_TYPES[ext]))]

    if ext == "docx":
        doc = DocxDocument(file_path)
        lines = [para.text for para in doc.paragraphs]
        for table in doc.tables:
            lines.extend(_table_rows(table))
        return [Document(page_content="\n".join(lines))]

    if ext == "pptx":
        slides = []
        for number, slide in enumerate(Presentation(file_path).slides, start=1):
            lines = []
            for shape in slide.shapes:
                if getattr(shape, "has_table", False) and shape.has_table:
                    lines.extend(_table_rows(shape.table))
                elif hasattr(shape, "text"):
                    lines.append(shape.text)
            slides.append(Document(page_content="\n".join(lines), metadata={"page": number}))
        return slides

    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return [Document(page_content=f.read())]


def split_documents(docs: list[Document], ext: str) -> list[dict]:
    """Chunks as {"text", "page"}; page is None for files without pages."""
    options = {"chunk_size": CHUNK_SIZE, "chunk_overlap": CHUNK_OVERLAP}
    if ext in CODE_LANGUAGES:
        splitter = RecursiveCharacterTextSplitter.from_language(CODE_LANGUAGES[ext], **options)
    else:
        splitter = RecursiveCharacterTextSplitter(separators=["\n\n", "\n", ". ", "。", " ", ""], **options)
    return [
        {"text": chunk.page_content, "page": chunk.metadata.get("page")}
        for chunk in splitter.split_documents(docs)
        if chunk.page_content.strip()
    ]
