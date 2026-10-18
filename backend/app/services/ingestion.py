"""Document ingestion: validation, private storage, text extraction, chunking and indexing.

File -> validation -> private storage -> extraction -> cleaning -> chunking -> embeddings -> ready.

A PDF is read page by page and a presentation slide by slide, so each chunk keeps the page it
came from. A Word file has no fixed pages, so its chunks keep the heading they stand under.
A spreadsheet is read row by row: every row is written out with its column names, and rows are
kept together in groups that name their sheet and row numbers.
Where no location is known the chunk has none, and none is invented later.
"""
import csv
import hashlib
import io
import logging
import re
import uuid
import zipfile
from pathlib import Path

import pymupdf
from datetime import date, datetime, time
from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openpyxl import load_workbook
from pptx import Presentation

from . import llm, retrieval
from ..core.config import settings
from ..core.database import documents, stored_files
from ..core.utils import now

log = logging.getLogger(__name__)

PARSER_VERSION = "1"
CHUNK_SIZE = 1200
CHUNK_OVERLAP = 200
# Rows of a spreadsheet kept in one chunk: below CHUNK_SIZE, so the splitter never cuts a row from its column names
ROW_GROUP_CHARS = 900

MIME_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
    "txt": "text/plain",
    "md": "text/markdown",
}
# The file every Word / PowerPoint / Excel archive must contain
_OFFICE_MARKER = {"docx": "word/document.xml", "pptx": "ppt/presentation.xml", "xlsx": "xl/workbook.xml"}
TABLE_TYPES = ("xlsx", "csv")


class RejectedFile(ValueError):
    """The upload is not acceptable; the message tells the user why."""


def file_extension(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def validate(filename: str, data: bytes) -> str:
    """Check extension, size and that the content really is what the extension says. Returns the extension."""
    ext = file_extension(filename)
    if ext not in MIME_TYPES:
        raise RejectedFile(f"Unsupported file type. Allowed: {', '.join(sorted(MIME_TYPES))}.")
    if not data:
        raise RejectedFile("The file is empty.")
    if len(data) > settings.MAX_UPLOAD_MB * 1024 * 1024:
        raise RejectedFile(f"The file is larger than {settings.MAX_UPLOAD_MB} MB.")

    if ext == "pdf":
        if not data.startswith(b"%PDF-"):
            raise RejectedFile("The file is not a valid PDF.")
    elif ext in _OFFICE_MARKER:
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = archive.namelist()
        except zipfile.BadZipFile:
            names = []
        if _OFFICE_MARKER[ext] not in names:
            raise RejectedFile(f"The file is not a valid .{ext} file.")
    else:
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            raise RejectedFile("The text file is not UTF-8 text.")
        if "\x00" in text:
            raise RejectedFile("The file is not a text file.")
    return ext


def store(organization_id: str, ext: str, data: bytes) -> str:
    """Save the original under a random name in private storage; returns its storage key."""
    key = f"{organization_id}/{uuid.uuid4().hex}.{ext}"
    path = settings.STORAGE_DIR / key
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    if settings.FILE_STORE == "database":
        stored_files.insert_one({"_id": key, "organization_id": organization_id, "data": data, "created_at": now()})
    return key


def stored_path(storage_key: str) -> Path:
    """Where the original is on disk. With FILE_STORE=database a file the disk has lost is put back first."""
    path = settings.STORAGE_DIR / storage_key
    if not path.is_file() and settings.FILE_STORE == "database":
        kept = stored_files.find_one({"_id": storage_key})
        if kept is not None:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(kept["data"])
    return path


def remove_stored(storage_key: str) -> None:
    (settings.STORAGE_DIR / storage_key).unlink(missing_ok=True)
    stored_files.delete_one({"_id": storage_key})


def reindex_missing() -> int:
    """Index again every ready document whose chunks are gone; returns how many there were.

    For a host that starts with an empty disk: the records in MongoDB still say "ready", but nothing
    can be found. A document whose original is gone as well becomes "failed" and says so.
    """
    lost = [d["_id"] for d in documents.find({"status": "ready"}, {"_id": 1}) if retrieval.count_chunks(str(d["_id"])) == 0]
    for document_id in lost:
        process_document(document_id)
    return len(lost)


def file_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _clean(text: str) -> str:
    """Tidy extracted text: re-join hyphenated line breaks and collapse stray whitespace."""
    text = text.replace("\x00", "")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _table_rows(table) -> list[str]:
    return [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]


def _load_docx(path: Path) -> list[Document]:
    """One Document per heading, so every chunk knows the section it belongs to."""
    doc = DocxDocument(str(path))
    sections: list[tuple[str | None, list[str]]] = [(None, [])]
    for para in doc.paragraphs:
        if para.style is not None and para.style.name.startswith("Heading") and para.text.strip():
            sections.append((para.text.strip(), [para.text]))
        else:
            sections[-1][1].append(para.text)
    for table in doc.tables:
        sections.append(("Tables", _table_rows(table)))
    return [
        Document(page_content=_clean("\n".join(lines)), metadata={"section": heading} if heading else {})
        for heading, lines in sections
    ]


def _cell(value) -> str:
    """A spreadsheet cell as it reads on screen: no time on a plain date, no ".0" on a whole number."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d" if value.time() == time(0) else "%Y-%m-%d %H:%M")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value)).strip()


def _sheet_documents(sheet: str | None, table: list[list]) -> list[Document]:
    """Groups of rows of one sheet. Row 1 holds the column names; every later row is written as "name: value"."""
    rows = [(number, [_cell(c) for c in row]) for number, row in enumerate(table, start=1)]
    rows = [(number, cells) for number, cells in rows if any(cells)]
    if not rows:
        return []
    names = [name or f"Column {i}" for i, name in enumerate(rows[0][1], start=1)]
    lines = [
        (number, f"Row {number}: " + "; ".join(
            f"{names[i] if i < len(names) else f'Column {i + 1}'}: {cell}" for i, cell in enumerate(cells) if cell
        ))
        for number, cells in rows[1:]
    ] or [(1, "Row 1: " + "; ".join(c for c in rows[0][1] if c))]  # a sheet with a single row has no column names

    where = f"sheet “{sheet}”, " if sheet else ""
    groups: list[list[tuple[int, str]]] = [[]]
    for line in lines:
        if groups[-1] and sum(len(text) + 1 for _, text in groups[-1]) + len(line[1]) > ROW_GROUP_CHARS:
            groups.append([])
        groups[-1].append(line)
    return [
        Document(
            page_content=(f"Sheet: {sheet}\n" if sheet else "") + "\n".join(text for _, text in group),
            metadata={"section": f"{where}rows {group[0][0]}-{group[-1][0]}" if len(group) > 1 else f"{where}row {group[0][0]}"},
        )
        for group in groups
    ]


def _load_table_file(path: Path, ext: str) -> list[Document]:
    if ext == "csv":
        return _sheet_documents(None, list(csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig")))))
    # data_only: the value a formula showed when the file was last saved, not the formula itself
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return [
            part for sheet in workbook.worksheets
            for part in _sheet_documents(sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)])
        ]
    finally:
        workbook.close()


def extract(path: Path, ext: str) -> list[Document]:
    """The text of a file as Documents whose metadata holds the page or section, where known."""
    if ext in TABLE_TYPES:
        return _load_table_file(path, ext)
    if ext == "pdf":
        with pymupdf.open(path) as pdf:
            return [
                Document(page_content=_clean(page.get_text()), metadata={"page": number})
                for number, page in enumerate(pdf, start=1)
            ]
    if ext == "docx":
        return _load_docx(path)
    if ext == "pptx":
        slides = []
        for number, slide in enumerate(Presentation(str(path)).slides, start=1):
            lines = []
            for shape in slide.shapes:
                if getattr(shape, "has_table", False) and shape.has_table:
                    lines.extend(_table_rows(shape.table))
                elif hasattr(shape, "text"):
                    lines.append(shape.text)
            slides.append(Document(page_content=_clean("\n".join(lines)), metadata={"page": number}))
        return slides
    return [Document(page_content=_clean(path.read_text(encoding="utf-8")))]


def split(docs: list[Document]) -> list[dict]:
    """Bounded, overlapping chunks cut on paragraph and sentence boundaries: {"text", "page", "section"}."""
    splitter = RecursiveCharacterTextSplitter(
        separators=["\n\n", "\n", ". ", " ", ""], chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP
    )
    return [
        {"text": c.page_content, "page": c.metadata.get("page"), "section": c.metadata.get("section")}
        for c in splitter.split_documents(docs)
        if c.page_content.strip()
    ]


def report_progress(document_id, done: int, total: int) -> None:
    """Record how many sections are embedded, for the status shown while a document processes.

    Only a display: a failed write is logged and never stops the processing itself.
    """
    try:
        documents.update_one(
            {"_id": document_id, "status": "processing"},
            {"$set": {"progress": {"stage": "indexing", "done": done, "total": total}}},
        )
    except Exception:
        log.exception("Could not record the progress of document %s", document_id)


def process_document(document_id) -> None:
    """Extract, chunk and index one stored document, and record the outcome on its record.

    Safe to run again after a failure: the old chunks are replaced, never added to. The document
    becomes "ready" only after every chunk is indexed; any error leaves it "failed" with the reason.
    """
    doc = documents.find_one_and_update(
        {"_id": document_id}, {"$set": {"status": "processing", "error": None, "progress": {"stage": "reading"}}}
    )
    if doc is None:
        return
    organization_id = doc["organization_id"]
    try:
        ext = file_extension(doc["filename"])
        path = stored_path(doc["storage_key"])
        if not path.is_file():
            raise ValueError("The original file is no longer in storage. Upload it again.")
        try:
            parts = extract(path, ext)
        except Exception:
            raise ValueError("The file could not be read (corrupt or password-protected).")
        if sum(len(p.page_content) for p in parts) < 10:
            raise ValueError(
                "The spreadsheet has no filled rows." if ext in TABLE_TYPES
                else "No readable text found. Scanned documents need OCR, which is not supported yet."
            )
        chunks = split(parts)
        if len(chunks) > settings.MAX_FILE_CHUNKS:
            raise ValueError(
                f"Too much text to index ({len(chunks)} sections, the limit is {settings.MAX_FILE_CHUNKS}). "
                "Split the file into smaller parts."
            )
        report_progress(document_id, 0, len(chunks))
        retrieval.index_document(
            organization_id, str(document_id), chunks, lambda done, total: report_progress(document_id, done, total)
        )
        documents.update_one({"_id": document_id}, {"$set": {
            "status": "ready",
            "progress": None,
            "chunk_count": len(chunks),
            "page_count": len(parts) if ext in ("pdf", "pptx") else None,
            "parser_version": PARSER_VERSION,
            "embedding_model": settings.EMBEDDING_MODEL,
            "processed_at": now(),
        }})
    except Exception as e:
        if isinstance(e, (ValueError, llm.LLMError)):
            reason = str(e)
        else:
            log.exception("Processing of document %s failed", document_id)
            reason = "Unexpected error while processing the file."
        try:
            retrieval.delete_document(organization_id, str(document_id))
        except Exception:
            log.exception("Could not remove chunks of failed document %s", document_id)
        documents.update_one(
            {"_id": document_id},
            {"$set": {"status": "failed", "error": reason, "progress": None, "chunk_count": 0, "processed_at": now()}},
        )
