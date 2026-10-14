import io
import threading
import time
from collections import defaultdict, deque
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from bson import ObjectId
from bson.errors import InvalidId
from fastapi import HTTPException, Response, status
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .config import settings

# Never sent to a browser
_PRIVATE_FIELDS = ("password_hash", "storage_key")


def now() -> datetime:
    return datetime.now(timezone.utc)


def oid(value: str) -> ObjectId:
    """An id from a URL or request body; a malformed one is simply a record that does not exist."""
    try:
        return ObjectId(value)
    except (InvalidId, TypeError):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Not found")


def clean(doc: dict) -> dict:
    """A database record as it is returned by the API."""
    out = {k: v for k, v in doc.items() if k not in _PRIVATE_FIELDS and k != "_id"}
    out["id"] = str(doc["_id"])
    return out


def image_type(data: bytes) -> str | None:
    """The media type of a PNG, JPEG or WebP picture, read from its first bytes; None for anything else.
    The browser's own claim is not trusted, and SVG is never accepted (it can carry scripts)."""
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if data.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    return None


def read_image(upload, max_kb: int) -> tuple[bytes, str]:
    """The bytes and media type of an uploaded picture, or a 413/415 the browser can show."""
    data = upload.file.read(max_kb * 1024 + 1)
    if len(data) > max_kb * 1024:
        raise HTTPException(status.HTTP_413_CONTENT_TOO_LARGE, f"The picture is larger than {max_kb} KB.")
    kind = image_type(data)
    if kind is None:
        raise HTTPException(status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, "Use a PNG, JPG or WebP picture.")
    return data, kind


def paginate(collection, query: dict, page: int, page_size: int, sort: list) -> dict:
    cursor = collection.find(query).sort(sort).skip((page - 1) * page_size).limit(page_size)
    return {
        "items": [clean(d) for d in cursor],
        "total": collection.count_documents(query),
        "page": page,
        "page_size": page_size,
    }


EXPORT_LIMIT = 20_000


def xlsx_download(filename: str, columns: list[tuple[str, str]], records) -> Response:
    """An Excel file of records for download. columns: (heading, field name); at most EXPORT_LIMIT rows."""
    zone = ZoneInfo(settings.LOCAL_TIMEZONE)
    book = Workbook()
    sheet = book.active
    sheet.append([heading for heading, _ in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for record in records:
        row = []
        for _, field in columns:
            value = record.get(field)
            if isinstance(value, datetime):
                # Excel has no time zones: the factory's local time, as everywhere else in the application
                value = value.astimezone(zone).replace(tzinfo=None, microsecond=0) if value.tzinfo else value
            elif isinstance(value, (list, dict)):
                value = ", ".join(str(v) for v in value) if isinstance(value, list) else str(value)
            row.append(value)
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # text typed by a user stays text, even when it starts with "="
    for index, (heading, _) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(12, min(48, len(heading) + 6))
    sheet.freeze_panes = "A2"
    data = io.BytesIO()
    book.save(data)
    return Response(
        data.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


class RateLimiter:
    """At most `limit` events per key in any 60 seconds. Kept in memory: enough for one server process."""

    def __init__(self, limit: int):
        self.limit = limit
        self._events: dict[str, deque] = defaultdict(deque)
        self._lock = threading.Lock()

    def check(self, key: str) -> None:
        cutoff = time.monotonic() - 60
        with self._lock:
            events = self._events[key]
            while events and events[0] < cutoff:
                events.popleft()
            if len(events) >= self.limit:
                raise HTTPException(
                    status.HTTP_429_TOO_MANY_REQUESTS,
                    "Too many requests. Wait a minute and try again.",
                    headers={"Retry-After": "60"},
                )
            events.append(time.monotonic())

    def reset(self) -> None:
        with self._lock:
            self._events.clear()
