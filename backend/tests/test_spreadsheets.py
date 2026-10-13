"""Excel and CSV files as documents to ask about, and the registers downloaded as Excel files."""
import io
from datetime import datetime

from openpyxl import Workbook, load_workbook

from backend.app.services import ingestion, retrieval

from .conftest import API

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook(sheets: dict[str, list[list]]) -> bytes:
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    data = io.BytesIO()
    book.save(data)
    return data.getvalue()


def read_export(response) -> list[list]:
    assert response.status_code == 200 and response.headers["content-type"] == XLSX, response.text
    return [list(row) for row in load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True)]


PARTS = workbook({
    "Spares": [
        ["Part number", "Description", "Machine", "Stock", "Reorder level", "Last checked"],
        ["HF-150", "Hydraulic oil filter", "IM-150", 4, 2, datetime(2026, 3, 1)],
        ["HB-3", "Heater band for barrel zone 3", "IM-150", 1, 2.0, datetime(2026, 3, 1, 14, 30)],
        [None, None, None, None, None, None],
        ["TC-J2", "Thermocouple type J", "IM-150", 6, 3, None],
    ],
    "PM schedule": [
        ["Task", "Interval (hours)"],
        ["Grease the toggle pins", 40],
        ["Replace the hydraulic oil filter", 1000],
    ],
    "Empty": [],
})


def test_excel_file_is_indexed_row_by_row_with_its_column_names(client, factory):
    doc = factory.document("spares.xlsx", PARTS)
    assert doc["status"] == "ready" and doc["mime_type"] == XLSX and doc["page_count"] is None
    assert doc["chunk_count"] == 2  # one group of rows per sheet; the empty sheet gives nothing

    found = retrieval.search(factory.organization_id, "heater band HB-3 stock and reorder level", 1)[0]
    assert found["document_id"] == doc["id"] and found["section"] == "sheet “Spares”, rows 2-5"
    # Every value stands next to its column name, a whole number has no ".0", a plain date has no time
    assert "Row 3: Part number: HB-3; Description: Heater band for barrel zone 3; Machine: IM-150; Stock: 1; Reorder level: 2; Last checked: 2026-03-01 14:30" in found["text"]
    assert "Row 2: Part number: HF-150" in found["text"] and "Last checked: 2026-03-01\n" in found["text"]
    assert "Row 5: Part number: TC-J2" in found["text"] and "Row 4" not in found["text"]
    schedule = retrieval.search(factory.organization_id, "how often to grease the toggle pins", 1)[0]
    assert schedule["section"] == "sheet “PM schedule”, rows 2-3" and "Interval (hours): 40" in schedule["text"]


def test_a_question_is_answered_from_a_spreadsheet_with_the_rows_as_source(client, factory, fake_llm):
    factory.document("spares.xlsx", PARTS)
    fake_llm.reply = "The stock of heater band HB-3 is 1 and its reorder level is 2 [1]."
    answer = client.post(f"{API}/chat", headers=factory.admin, json={"question": "What is the stock of heater band HB-3?"}).json()
    assert answer["insufficient_evidence"] is False
    assert answer["citations"][0]["filename"] == "spares.xlsx" and answer["citations"][0]["section"] == "sheet “Spares”, rows 2-5"
    assert "Stock: 1" in fake_llm.calls[0]["prompt"]


def test_csv_file_and_long_sheets_are_split_between_rows(client, factory):
    doc = factory.document("limits.csv", "﻿Parameter,Limit\nOil temperature,55 C\nWater pressure,3 bar\n".encode())
    assert doc["status"] == "ready" and doc["mime_type"] == "text/csv" and doc["chunk_count"] == 1
    found = retrieval.search(factory.organization_id, "oil temperature limit", 1)[0]
    assert found["section"] == "rows 2-3" and "Row 2: Parameter: Oil temperature; Limit: 55 C" in found["text"]

    rows = [["Code", "Meaning"]] + [[f"A-{n}", f"Alarm number {n}: check the sensor of zone {n} and its cable"] for n in range(100, 160)]
    long = factory.document("alarms.xlsx", workbook({"Alarms": rows}))
    assert long["status"] == "ready" and long["chunk_count"] > 3
    parts = ingestion.extract(ingestion.stored_path(_storage_key(long["id"])), "xlsx")
    assert all(len(p.page_content) <= ingestion.CHUNK_SIZE for p in parts)
    # No row is cut in two, and none is lost or repeated
    lines = [line for p in parts for line in p.page_content.splitlines() if line.startswith("Row ")]
    assert [line.split(":")[0] for line in lines] == [f"Row {n}" for n in range(2, 62)]
    assert parts[0].metadata["section"].startswith("sheet “Alarms”, rows 2-")


def _storage_key(document_id: str) -> str:
    from backend.app.core import database
    from backend.app.core.utils import oid

    return database.documents.find_one({"_id": oid(document_id)})["storage_key"]


def test_bad_spreadsheets_are_refused_or_fail_with_a_reason(client, factory):
    assert factory.upload("fake.xlsx", b"not an excel file").status_code == 422
    assert factory.upload("old.xls", b"anything").status_code == 422
    empty = factory.document("empty.xlsx", workbook({"Sheet1": []}))
    assert empty["status"] == "failed" and empty["error"] == "The spreadsheet has no filled rows."


def test_registers_are_exported_to_excel(client, factory):
    asset = factory.asset("IM-150")
    factory.asset("CNC-01")
    factory.breakdown(asset["id"], error_code="A-207", downtime_minutes=90, root_cause="=1+1")
    factory.breakdown(asset["id"], error_code="A-315", symptoms="Oil temperature high")

    table = read_export(client.get(f"{API}/breakdowns/export", headers=factory.admin))
    assert table[0][:6] == ["Record", "Asset", "Reported", "Symptoms", "Error code", "Downtime (min)"]
    assert len(table) == 3 and {row[4] for row in table[1:]} == {"A-207", "A-315"}
    row = next(r for r in table[1:] if r[4] == "A-207")
    assert row[1] == "IM-150" and row[5] == 90 and row[2] == datetime(2026, 3, 1, 8, 0)  # LOCAL_TIMEZONE is UTC in tests
    assert row[6] == "=1+1"  # kept as text, not turned into a formula
    # The same filters as the list
    one = read_export(client.get(f"{API}/breakdowns/export", headers=factory.admin, params={"error_code": "a-315"}))
    assert len(one) == 2 and one[1][3] == "Oil temperature high"

    assets = read_export(client.get(f"{API}/assets/export", headers=factory.admin))
    assert [r[0] for r in assets] == ["Asset code", "CNC-01", "IM-150"]

    client.post(f"{API}/work-orders", headers=factory.admin, json={
        "asset_id": asset["id"], "title": "Replace heater band", "checklist": [{"step": "Switch off"}, {"step": "Replace HB-3"}],
    })
    orders = read_export(client.get(f"{API}/work-orders/export", headers=factory.admin, params={"status": "draft"}))
    header = orders[0]
    assert orders[1][header.index("Title")] == "Replace heater band" and orders[1][header.index("Asset")] == "IM-150"
    assert orders[1][header.index("Steps")] == "1. Switch off\n2. Replace HB-3"
    assert len(read_export(client.get(f"{API}/work-orders/export", headers=factory.admin, params={"status": "approved"}))) == 1

    events = client.get(f"{API}/audit", headers=factory.admin).json()["items"]
    assert {"breakdown.export", "asset.export", "work_order.export"} <= {e["action"] for e in events}


def test_export_follows_access_and_stays_inside_the_factory(client, factory, other_factory):
    factory.breakdown(factory.asset("IM-150")["id"])
    assert len(read_export(client.get(f"{API}/breakdowns/export", headers=other_factory.admin))) == 1  # only the heading
    created = client.post(f"{API}/users", headers=factory.admin, json={
        "name": "Ask only", "email": "ask@alpha.example.com", "password": "test-password-1", "role": "viewer",
        "permissions": ["chat.ask"],
    })
    assert created.status_code == 201
    asker = factory.login("ask@alpha.example.com")
    for path in ("breakdowns", "assets", "work-orders"):
        assert client.get(f"{API}/{path}/export", headers=asker).status_code == 403, path
